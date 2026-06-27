import os
import json
import traceback
from datetime import datetime, timedelta
from sqlalchemy import func
from db.database import SessionLocal
from db.models import (Requests, IssueType, Category, Status, 
                       RequestDetails, Materials, CriticalityLevels)
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BUDGET_FILE = os.path.join(BASE_DIR, "budget_config.json")

UKR_MONTHS = {
    1: "січень", 2: "лютий", 3: "березень", 4: "квітень",
    5: "травень", 6: "червень", 7: "липень", 8: "серпень",
    9: "вересень", 10: "жовтень", 11: "листопад", 12: "грудень"
}

def _load_budget_data():
    """Зчитування конфігураційного файлу розподілу бюджетних коштів."""
    if os.path.exists(BUDGET_FILE):
        try:
            with open(BUDGET_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except: pass
    return {}

def _save_budget_data(data):
    """Збереження оновлених лімітів бюджету у конфігураційний файл."""
    with open(BUDGET_FILE, 'w', encoding='utf-8') as f:
        json.dump(data, f, indent=4)

def get_budget(year, month):
    data = _load_budget_data()
    y_str, m_str = str(year), str(month).zfill(2)
    return data.get(y_str, {}).get(m_str, 0.0) 

def get_yearly_budget(year):
    data = _load_budget_data()
    y_str = str(year)
    if y_str in data: return sum(data[y_str].values())
    return 0.0

def set_budget(year, month, amount):
    data = _load_budget_data()
    y_str, m_str = str(year), str(month).zfill(2)
    if y_str not in data: data[y_str] = {}
    data[y_str][m_str] = float(amount)
    _save_budget_data(data)

def get_efficiency_data(start_date, end_date, category_filter):
    """Формування статистичних даних завантаженості та ефективності категорій інцидентів."""
    with SessionLocal() as db:
        try:
            all_categories = db.query(Category).all()
            cat_names = [c.category_name for c in all_categories]
            chart_data_dict = {name: 0 for name in cat_names}

            query = db.query(Requests).join(Status).join(IssueType).join(Category)
            
            if start_date: query = query.filter(Requests.request_date >= start_date)
            if end_date:
                end_date_inclusive = end_date + timedelta(days=1)
                query = query.filter(Requests.request_date < end_date_inclusive)
                
            if category_filter and category_filter != "Всі категорії":
                query = query.filter(Category.category_name.ilike(f"%{category_filter}%"))
                
            requests = query.all()
            counts = {"new": 0, "in_progress": 0, "completed": 0, "cancelled": 0}
            table_data = []
            
            for r in requests:
                status_name = r.status.status_name.lower() if r.status else ""
                cat_name = r.issue_type.category.category_name if r.issue_type and r.issue_type.category else ""
                if cat_name in chart_data_dict: chart_data_dict[cat_name] += 1
                
                if "нова" in status_name or "новий" in status_name: counts["new"] += 1
                elif "виконано" in status_name: counts["completed"] += 1
                elif "скасовано" in status_name: counts["cancelled"] += 1
                else: counts["in_progress"] += 1

                crew_num = f"Бригада №{r.crew.crew_number}" if r.crew else "Не призначено"
                table_data.append([
                    str(r.id_request), r.request_date.strftime("%d.%m.%Y %H:%M") if r.request_date else "",
                    cat_name, crew_num, r.criticality.level_name if r.criticality else "Не визначено", r.status.status_name
                ])
            table_data.sort(key=lambda x: x[1], reverse=True)
            return {
                "categories": cat_names, "counts": counts, "chart_data": list(chart_data_dict.values()),
                "chart_labels": list(chart_data_dict.keys()), "table_data": table_data[:20]
            }
        except Exception: return {"categories": [], "counts": {"new":0, "in_progress":0, "completed":0, "cancelled":0}, "chart_data": [], "chart_labels": [], "table_data": []}

def get_budget_data(year, month):
    """Агрегація фінансових витрат комунального підприємства за обраний період часу."""
    with SessionLocal() as db:
        try:
            total_budget = 0.0
            report_title = ""
            file_suffix = ""
            
            query = db.query(RequestDetails).join(Requests).outerjoin(Materials)
            
            if month == -1:
                report_title = "за весь час"
                file_suffix = "весь_час"
                all_data = _load_budget_data()
                if all_data:
                    for y, m_data in all_data.items():
                        for m_str, val in m_data.items():
                            try: total_budget += float(val)
                            except: pass
            else:
                filter_date = func.coalesce(Requests.completion_date, Requests.request_date)
                if month == 0: 
                    start_date = datetime(year, 1, 1)
                    end_date = datetime(year + 1, 1, 1)
                    total_budget = get_yearly_budget(year)
                    report_title = f"за {year} рік"
                    file_suffix = f"{year}рік"
                else:
                    start_date = datetime(year, month, 1)
                    end_date = datetime(year + 1, 1, 1) if month == 12 else datetime(year, month + 1, 1)
                    total_budget = get_budget(year, month)
                    report_title = f"за {UKR_MONTHS[month]} {year} року"
                    file_suffix = f"{UKR_MONTHS[month]}_{year}рік"
                
                query = query.filter(filter_date >= start_date, filter_date < end_date)
            
            details = query.all()
            total_spent = 0.0
            table_data = []
            
            for d in details:
                try:
                    cost = float(d.total_cost) if d.total_cost else 0.0
                    total_spent += cost
                    
                    qty = float(d.quantity) if d.quantity else 0.0
                    material_name = d.material.material_name if d.material else "Видалений матеріал"
                    price = f"{d.material.price} ₴" if d.material else "0 ₴"
                    unit = d.material.unit if d.material else "шт"
                    
                    table_data.append([
                        f"Заявка #{d.request_id}", 
                        material_name,
                        price, 
                        f"{qty} {unit}",
                        f"{cost:.2f} ₴"
                    ])
                except Exception as ex:
                    print(f"Помилка обробки рядка витрат ID {d.id_detail}: {ex}")
                
            return {
                "period_label": report_title, "file_suffix": file_suffix, "total_budget": total_budget,
                "spent": total_spent, "remaining": total_budget - total_spent, "table_data": table_data
            }
        except Exception as e:
            traceback.print_exc()
            return {"period_label": "Помилка", "file_suffix": "помилка", "total_budget": 0, "spent": 0, "remaining": 0, "table_data": []}

def get_map_statistics():
    with SessionLocal() as db:
        stats = []
        try:
            query = db.query(Category.category_name, CriticalityLevels.level_name, func.count(Requests.id_request).label('count'))\
             .join(IssueType, Requests.issue_type_id == IssueType.id_issue_type).join(Category, IssueType.category_id == Category.id_category)\
             .join(CriticalityLevels, Requests.criticality_id == CriticalityLevels.id_criticality).join(Status, Requests.status_id == Status.id_status)\
             .filter(~Status.status_name.ilike("%Виконано%"), ~Status.status_name.ilike("%Скасовано%"))\
             .group_by(Category.category_name, CriticalityLevels.level_name).all()
            for row in query: stats.append({"category": row.category_name, "criticality": row.level_name, "count": row.count})
        except Exception: pass
        return stats

def generate_word_report(path, budget_data):
    """Генерація текстового фінансового звіту у форматі Microsoft Word (*.docx)."""
    doc = Document()
    title = doc.add_heading('Звіт про витрати комунального підприємства', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_paragraph(f"Дата формування: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
    doc.add_heading(f"Фінансові показники ({budget_data['period_label']})", level=1)
    p = doc.add_paragraph()
    p.add_run("Загальний бюджет: ").bold = True
    p.add_run(f"{budget_data['total_budget']:,.2f} грн\n")
    p.add_run("Вже витрачено: ").bold = True
    p.add_run(f"{budget_data['spent']:,.2f} грн\n")
    p.add_run("Залишок фонду: ").bold = True
    p.add_run(f"{budget_data['remaining']:,.2f} грн")
    
    doc.add_heading('Деталізація витрат', level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    headers = ["Заявка", "Матеріал", "Ціна", "Кількість", "Сума"]
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].bold = True
    for row_data in budget_data['table_data']:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data): row_cells[i].text = str(val)
    doc.save(path)

def generate_excel_report(path, budget_data):
    """Генерація табличного фінансового документа у форматі Microsoft Excel (*.xlsx)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Витрати"
    ws.append([f"Звіт про витрати КП ({budget_data['period_label']})"])
    ws.merge_cells('A1:E1')
    ws['A1'].font = Font(size=14, bold=True)
    ws['A1'].alignment = Alignment(horizontal="center")
    ws.append([])
    ws.append(["Загальний бюджет:", f"{budget_data['total_budget']:,.2f} грн"])
    ws.append(["Вже витрачено:", f"{budget_data['spent']:,.2f} грн"])
    ws.append(["Залишок фонду:", f"{budget_data['remaining']:,.2f} грн"])
    for i in range(2, 5): ws.cell(row=i, column=1).font = Font(bold=True)
    ws.append([])
    ws.append(["Заявка", "Матеріал", "Ціна", "Кількість", "Сума"])
    for col in range(1, 6): ws.cell(row=7, column=col).font = Font(bold=True)
    for row_data in budget_data['table_data']: ws.append(row_data)
    wb.save(path)

def get_status_duration_statistics():
    """Аналізує таблицю status_history для розрахунку середнього часу
      перебування заявок у кожному зі статусів (у годинах та хвилинах)."""
    from db.models import StatusHistory, Status

    with SessionLocal() as db:
        try:
            histories = db.query(StatusHistory).join(Status) \
                .order_by(StatusHistory.request_id, StatusHistory.change_date.asc()).all()

            durations = {}

            for i in range(len(histories) - 1):
                current_action = histories[i]
                next_action = histories[i + 1]

                if current_action.request_id == next_action.request_id:
                    diff_seconds = (next_action.change_date - current_action.change_date).total_seconds()
                    status_name = current_action.status.status_name

                    if status_name not in durations:
                        durations[status_name] = []
                    durations[status_name].append(diff_seconds)

            report_data = []
            for status_name, sec_list in durations.items():
                avg_seconds = sum(sec_list) / len(sec_list) if sec_list else 0

                hours = int(avg_seconds // 3600)
                minutes = int((avg_seconds % 3600) // 60)
                if hours == 0 and minutes == 0:
                    time_str = f"{int(avg_seconds)} сек (Миттєво)"
                else:
                    time_str = f"{hours} год. {minutes} хв."

                report_data.append({
                    "status_name": status_name,
                    "avg_seconds": avg_seconds,
                    "time_display": time_str,
                    "transitions_count": len(sec_list)
                })

            report_data.sort(key=lambda x: x["avg_seconds"], reverse=True)
            return report_data
        except Exception as e:
            print(f"Помилка розрахунку SLA метрик: {e}")
            return []